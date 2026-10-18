# -*- coding: utf-8 -*-
"""
生成《专业课程设计II报告》Word 文档。

分工：
  · report_lib.py      排版基元（字体、缩进、图注、表格、代码块）
  · report_content.py  正文内容（模板的七个部分）
  · 本文件             装配封面、评分表与正文，并设置页面属性

用法：python scripts/build_report.py
输出：docs/专业课程设计II报告_学号.docx

生成后请自行完成两件事：
  1. 替换封面上的【专业】【姓名】【班级】【学号】占位符；
  2. 通读一遍，把与人称、分工相关的表述调整为你自己的实际情况。
"""

import os
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

import report_content as C
from report_lib import (EN_FONT, OUT_PATH, CN_FONT, SZ_S4, build_cover,
                        build_score_sheet, page_break, set_font)

# 允许用命令行参数覆盖输出路径：
#   python scripts/build_report.py D:/.../专业课程设计II报告_终版.docx
# 用途：目标文件被 Word 或资源管理器预览窗格占用时，改写到另一个文件名。
OUTPUT_PATH = sys.argv[1] if len(sys.argv) > 1 else None


def setup_document():
    doc = Document()

    # 页面：A4，上下 2.54cm、左右 3.17cm（Word 默认值，符合中文报告习惯）
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)

    # 默认样式：中文宋体、英文 Times New Roman、小四
    style = doc.styles['Normal']
    style.font.name = EN_FONT
    style.font.size = SZ_S4
    style.element.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT)

    return doc


def add_page_numbers(doc):
    """在页脚居中插入页码。报告要求不少于 15 页，页码便于评审时翻阅。"""
    from docx.oxml import OxmlElement

    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = p.add_run()
        set_font(run, size=Pt(10.5))

        fld_begin = OxmlElement('w:fldChar')
        fld_begin.set(qn('w:fldCharType'), 'begin')
        instr = OxmlElement('w:instrText')
        instr.set(qn('xml:space'), 'preserve')
        instr.text = 'PAGE'
        fld_end = OxmlElement('w:fldChar')
        fld_end.set(qn('w:fldCharType'), 'end')

        run._r.append(fld_begin)
        run._r.append(instr)
        run._r.append(fld_end)


def main():
    doc = setup_document()
    add_page_numbers(doc)

    build_cover(doc)
    build_score_sheet(doc)

    # 正文（各节内部通过 figure() 自动连续编号图号）
    C.build_section1(doc)
    C.build_section2(doc)
    C.build_section3(doc)
    C.build_section4(doc)
    C.build_section5(doc)
    C.build_section6(doc)
    C.build_section7(doc)

    # 目标文件可能正被 Word 打开、或被资源管理器的预览窗格占用。
    # 这种情况下直接保存会抛 PermissionError 且提示不明确，
    # 因此退而写入带序号的文件名，并明确告诉使用者原因与处理办法。
    target = OUTPUT_PATH if OUTPUT_PATH else OUT_PATH
    os.makedirs(os.path.dirname(target), exist_ok=True)

    saved_path = target
    try:
        doc.save(target)
    except PermissionError:
        stem, ext = os.path.splitext(target)
        for i in range(1, 10):
            candidate = '%s_v%d%s' % (stem, i, ext)
            try:
                doc.save(candidate)
                saved_path = candidate
                break
            except PermissionError:
                continue
        else:
            raise SystemExit(
                '目标文件与多个备选文件名均被占用，无法写入。\n'
                '请关闭正在打开该文档的 Word 窗口，并关闭资源管理器的「预览窗格」，'
                '然后重新执行本脚本。')

        print('!' * 62)
        print('原报告文件被占用，无法覆盖：')
        print('   %s' % target)
        print('   常见原因：该文档正在 Word 中打开，或资源管理器开启了预览窗格。')
        print('   本次已改为输出到：')
        print('   %s' % saved_path)
        print('!' * 62)
        print()

    size_kb = os.path.getsize(saved_path) / 1024
    print('报告已生成：%s' % saved_path)
    print('文件大小：%.0f KB' % size_kb)
    print()
    print('请完成以下两件事后再提交：')
    print('  1. 在 Word 中替换封面上的【专业】【姓名】【班级】【学号】占位符；')
    print('  2. 通读全文，把与人称、分工相关的表述调整为你自己的实际情况。')


if __name__ == '__main__':
    main()
